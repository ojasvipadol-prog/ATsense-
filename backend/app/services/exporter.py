import os
from io import BytesIO
from typing import Dict, Any
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def generate_pdf_resume(content: Dict[str, Any], template: str = "classic") -> bytes:
    """Generates an authentic, beautifully-formatted PDF resume according to the selected template."""
    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        rightMargin=0.5 * inch,
        leftMargin=0.5 * inch,
        topMargin=0.5 * inch,
        bottomMargin=0.5 * inch
    )
    
    styles = getSampleStyleSheet()
    
    # Custom color themes per template
    if template == "modern":
        primary_color = colors.HexColor("#1e3a8a")  # Deep Blue
        accent_color = colors.HexColor("#3b82f6")   # Royal Blue
        divider_color = colors.HexColor("#93c5fd")
    elif template == "technical":
        primary_color = colors.HexColor("#0f172a")  # Slate 900
        accent_color = colors.HexColor("#0284c7")   # Sky Blue
        divider_color = colors.HexColor("#cbd5e1")
    elif template == "fresher":
        primary_color = colors.HexColor("#1e293b")  # Dark Slate
        accent_color = colors.HexColor("#10b981")   # Emerald
        divider_color = colors.HexColor("#a7f3d0")
    else:  # classic ATS
        primary_color = colors.HexColor("#111827")  # Classic Black/Charcoal
        accent_color = colors.HexColor("#374151")   # Gray 700
        divider_color = colors.HexColor("#9ca3af")

    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=20,
        leading=24,
        textColor=primary_color,
        alignment=1 if template in ["classic", "fresher"] else 0, # Centered or Left
        spaceAfter=4
    )
    
    subtitle_style = ParagraphStyle(
        'DocSubtitle',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9,
        leading=12,
        textColor=colors.HexColor("#4b5563"),
        alignment=1 if template in ["classic", "fresher"] else 0,
        spaceAfter=10
    )
    
    section_head_style = ParagraphStyle(
        'SectionHead',
        parent=styles['Heading2'],
        fontName='Helvetica-Bold',
        fontSize=12,
        leading=15,
        textColor=primary_color,
        spaceBefore=8,
        spaceAfter=4
    )
    
    item_title_style = ParagraphStyle(
        'ItemTitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=10,
        leading=13,
        textColor=colors.HexColor("#1f2937")
    )
    
    body_style = ParagraphStyle(
        'Body',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9,
        leading=12,
        textColor=colors.HexColor("#374151")
    )
    
    bullet_style = ParagraphStyle(
        'Bullet',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8.5,
        leading=12,
        leftIndent=12,
        textColor=colors.HexColor("#374151")
    )
    
    story = []
    
    # 1. Header & Personal Info
    personal = content.get("personal", {})
    full_name = personal.get("fullName", "Candidate Name")
    story.append(Paragraph(full_name, title_style))
    
    contact_parts = []
    if personal.get("email"): contact_parts.append(personal.get("email"))
    if personal.get("phone"): contact_parts.append(personal.get("phone"))
    if personal.get("location"): contact_parts.append(personal.get("location"))
    if personal.get("linkedin"): contact_parts.append(personal.get("linkedin"))
    if personal.get("github"): contact_parts.append(personal.get("github"))
    if personal.get("portfolio"): contact_parts.append(personal.get("portfolio"))
    
    story.append(Paragraph(" • ".join(contact_parts), subtitle_style))
    story.append(HRFlowable(width="100%", thickness=1, color=divider_color, spaceBefore=0, spaceAfter=8))
    
    # 2. Professional Summary
    summary = personal.get("summary")
    if summary:
        story.append(Paragraph("PROFESSIONAL SUMMARY", section_head_style))
        story.append(HRFlowable(width="100%", thickness=0.5, color=divider_color, spaceBefore=0, spaceAfter=4))
        story.append(Paragraph(summary, body_style))
        story.append(Spacer(1, 6))
        
    # Helper for adding sections based on template order preference
    def add_skills():
        skills = content.get("skills", {})
        if skills:
            story.append(Paragraph("TECHNICAL SKILLS", section_head_style))
            story.append(HRFlowable(width="100%", thickness=0.5, color=divider_color, spaceBefore=0, spaceAfter=4))
            for cat, s_list in skills.items():
                if isinstance(s_list, list) and s_list:
                    cat_text = f"<b>{cat}:</b> {', '.join(s_list)}"
                    story.append(Paragraph(cat_text, body_style))
            story.append(Spacer(1, 6))
            
    def add_education():
        education = content.get("education", [])
        if education:
            story.append(Paragraph("EDUCATION", section_head_style))
            story.append(HRFlowable(width="100%", thickness=0.5, color=divider_color, spaceBefore=0, spaceAfter=4))
            for edu in education:
                degree = edu.get("degree", "")
                univ = edu.get("university") or edu.get("college", "")
                year = edu.get("year", "")
                cgpa = edu.get("cgpa", "")
                header_line = f"<b>{degree}</b> — {univ}"
                right_line = f"{year} {f'| CGPA: {cgpa}' if cgpa else ''}"
                story.append(Paragraph(f"{header_line} <font color='#6b7280'>({right_line})</font>", body_style))
            story.append(Spacer(1, 6))

    def add_experience():
        experience = content.get("experience", [])
        if experience:
            story.append(Paragraph("PROFESSIONAL EXPERIENCE", section_head_style))
            story.append(HRFlowable(width="100%", thickness=0.5, color=divider_color, spaceBefore=0, spaceAfter=4))
            for exp in experience:
                role = exp.get("position", "")
                comp = exp.get("company", "")
                dur = exp.get("duration", "")
                story.append(Paragraph(f"<b>{role}</b> | {comp} <font color='#6b7280'>({dur})</font>", item_title_style))
                resp = exp.get("responsibilities", "")
                ach = exp.get("achievements", "")
                if resp:
                    story.append(Paragraph(f"• {resp}", bullet_style))
                if ach:
                    story.append(Paragraph(f"• Key Achievement: {ach}", bullet_style))
                story.append(Spacer(1, 4))
            story.append(Spacer(1, 4))

    def add_projects():
        projects = content.get("projects", [])
        if projects:
            story.append(Paragraph("KEY PROJECTS", section_head_style))
            story.append(HRFlowable(width="100%", thickness=0.5, color=divider_color, spaceBefore=0, spaceAfter=4))
            for proj in projects:
                name = proj.get("name", "")
                tech = proj.get("technologies", "")
                desc = proj.get("description", "")
                resp = proj.get("responsibilities", "")
                results = proj.get("results", "")
                
                title_line = f"<b>{name}</b>" + (f" | <i>{tech}</i>" if tech else "")
                story.append(Paragraph(title_line, item_title_style))
                if desc:
                    story.append(Paragraph(f"• {desc}", bullet_style))
                if resp:
                    story.append(Paragraph(f"• {resp}", bullet_style))
                if results:
                    story.append(Paragraph(f"• Impact & Results: {results}", bullet_style))
                story.append(Spacer(1, 4))
            story.append(Spacer(1, 4))

    def add_certifications():
        certs = content.get("certifications", [])
        if certs:
            story.append(Paragraph("CERTIFICATIONS", section_head_style))
            story.append(HRFlowable(width="100%", thickness=0.5, color=divider_color, spaceBefore=0, spaceAfter=4))
            for cert in certs:
                story.append(Paragraph(f"• {cert}", bullet_style))
            story.append(Spacer(1, 6))

    def add_achievements():
        achievements = content.get("achievements", [])
        if achievements:
            story.append(Paragraph("HONORS & ACHIEVEMENTS", section_head_style))
            story.append(HRFlowable(width="100%", thickness=0.5, color=divider_color, spaceBefore=0, spaceAfter=4))
            for ach in achievements:
                story.append(Paragraph(f"• {ach}", bullet_style))
            story.append(Spacer(1, 6))

    # Arrange sections based on template archetype
    if template == "fresher":
        # Education and projects prioritized first
        add_education()
        add_skills()
        add_projects()
        add_experience()
        add_certifications()
        add_achievements()
    elif template == "technical":
        # Technical skills and projects first
        add_skills()
        add_projects()
        add_experience()
        add_education()
        add_certifications()
        add_achievements()
    else:
        # Standard ATS Classic / Modern Professional
        add_experience()
        add_projects()
        add_skills()
        add_education()
        add_certifications()
        add_achievements()
        
    doc.build(story)
    buffer.seek(0)
    return buffer.getvalue()

def generate_docx_resume(content: Dict[str, Any], template: str = "classic") -> bytes:
    """Generates an authentic Microsoft Word (.docx) resume with structured headings and styles."""
    doc = docx.Document()
    
    # Page Margins 0.6 inch
    sections = doc.sections
    for s in sections:
        s.top_margin = Inches(0.6)
        s.bottom_margin = Inches(0.6)
        s.left_margin = Inches(0.6)
        s.right_margin = Inches(0.6)
        
    personal = content.get("personal", {})
    
    # Title / Name
    title_p = doc.add_paragraph()
    if template in ["classic", "fresher"]:
        title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run(personal.get("fullName", "Candidate Name"))
    title_run.font.size = Pt(20)
    title_run.font.bold = True
    title_run.font.name = 'Calibri'
    title_run.font.color.rgb = RGBColor(30, 58, 138) if template == "modern" else RGBColor(17, 24, 39)
    
    # Subtitle / Contact
    contact_parts = []
    if personal.get("email"): contact_parts.append(personal.get("email"))
    if personal.get("phone"): contact_parts.append(personal.get("phone"))
    if personal.get("location"): contact_parts.append(personal.get("location"))
    if personal.get("linkedin"): contact_parts.append(personal.get("linkedin"))
    if personal.get("github"): contact_parts.append(personal.get("github"))
    
    contact_p = doc.add_paragraph()
    if template in ["classic", "fresher"]:
        contact_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact_run = contact_p.add_run(" | ".join(contact_parts))
    contact_run.font.size = Pt(9.5)
    contact_run.font.color.rgb = RGBColor(107, 114, 128)
    
    def add_heading(text: str):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(8)
        h.paragraph_format.space_after = Pt(2)
        r = h.add_run(text.upper())
        r.font.size = Pt(11)
        r.font.bold = True
        r.font.name = 'Calibri'
        r.font.color.rgb = RGBColor(30, 58, 138) if template == "modern" else RGBColor(31, 41, 55)
        
    # Summary
    if personal.get("summary"):
        add_heading("Professional Summary")
        p = doc.add_paragraph(personal.get("summary"))
        p.paragraph_format.space_after = Pt(4)
        
    # Skills
    skills = content.get("skills", {})
    if skills:
        add_heading("Technical Skills")
        for cat, s_list in skills.items():
            if isinstance(s_list, list) and s_list:
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(2)
                bold_cat = p.add_run(f"{cat}: ")
                bold_cat.font.bold = True
                p.add_run(", ".join(s_list))
                
    # Projects
    projects = content.get("projects", [])
    if projects:
        add_heading("Key Projects")
        for proj in projects:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(1)
            name_run = p.add_run(proj.get("name", ""))
            name_run.font.bold = True
            if proj.get("technologies"):
                tech_run = p.add_run(f" | {proj.get('technologies')}")
                tech_run.font.italic = True
            if proj.get("description"):
                doc.add_paragraph(proj.get("description"), style='List Bullet')
            if proj.get("responsibilities"):
                doc.add_paragraph(proj.get("responsibilities"), style='List Bullet')
            if proj.get("results"):
                doc.add_paragraph(f"Results: {proj.get('results')}", style='List Bullet')

    # Experience
    experience = content.get("experience", [])
    if experience:
        add_heading("Professional Experience")
        for exp in experience:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(1)
            role_run = p.add_run(exp.get("position", ""))
            role_run.font.bold = True
            comp_run = p.add_run(f" — {exp.get('company', '')} ({exp.get('duration', '')})")
            if exp.get("responsibilities"):
                doc.add_paragraph(exp.get("responsibilities"), style='List Bullet')
            if exp.get("achievements"):
                doc.add_paragraph(f"Impact: {exp.get('achievements')}", style='List Bullet')

    # Education
    education = content.get("education", [])
    if education:
        add_heading("Education")
        for edu in education:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            deg_run = p.add_run(edu.get("degree", ""))
            deg_run.font.bold = True
            univ = edu.get("university") or edu.get("college", "")
            p.add_run(f", {univ} ({edu.get('year', '')})")
            if edu.get("cgpa"):
                p.add_run(f" | CGPA: {edu.get('cgpa')}")

    # Certifications
    certs = content.get("certifications", [])
    if certs:
        add_heading("Certifications")
        for c in certs:
            doc.add_paragraph(c, style='List Bullet')

    # Achievements
    achievements = content.get("achievements", [])
    if achievements:
        add_heading("Honors & Achievements")
        for a in achievements:
            doc.add_paragraph(a, style='List Bullet')

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()
